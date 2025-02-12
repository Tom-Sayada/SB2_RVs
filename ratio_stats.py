import os
import re
import pandas as pd
import numpy as np
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def process_simulation_folder(base_path):
    for parent_root, parent_dirs, _ in os.walk(base_path):
        # Identify parent folders containing simulation_x folders
        simulation_folders = [d for d in parent_dirs if d.startswith('simulation_')]
        if not simulation_folders:
            continue

        parent_folder = os.path.basename(parent_root)

        # Extract K1, K2 values from the parent folder name using regex
        try:
            k1_match = re.search(r'K_1_(\d+\.\d+)', parent_folder)
            k2_match = re.search(r'K_2_(\d+\.\d+)', parent_folder)

            if not all([k1_match, k2_match]):
                raise ValueError("Missing K1 or K2 values in folder name")

            K1 = float(k1_match.group(1))
            K2 = float(k2_match.group(1))
            ratio_real = K1 / K2
        except Exception as e:
            print(f"Skipping folder due to naming convention issue: {parent_folder}, Error: {e}")
            continue

        # Initialize data storage for all simulations within the parent folder
        data = []

        for simulation_folder in simulation_folders:
            simulation_path = os.path.join(parent_root, simulation_folder)
            sim_data = {'simulation': simulation_folder, 'ratio_real': ratio_real}

            # Process voigt fit
            voigt_folder = os.path.join(simulation_path, 'standard_sym_voigt_weighted_nobaseline_fit_results')
            if os.path.exists(voigt_folder):
                voigt_file = os.path.join(voigt_folder, 'fit_results.xlsx')
                if os.path.exists(voigt_file):
                    try:
                        df_chi = pd.read_excel(voigt_file, sheet_name='Chi_Square_Statistics')
                        global_row = df_chi[df_chi.iloc[:, 0] == 'GLOBAL']
                        if not global_row.empty:
                            sim_data.update({
                                'ratio_calc_voigt': global_row['ratio'].values[0],
                                'ratio_calc_voigt_err': global_row['ratio_err'].values[0],
                                'real_calc_diff_over_err_voigt': abs(global_row['ratio'].values[0] - ratio_real) / global_row['ratio_err'].values[0]
                            })
                    except Exception as e:
                        print(f"Error reading voigt file {voigt_file}: {e}")

            # Process gaussian fit
            gaussian_folder = os.path.join(simulation_path, 'standard_sym_gaussian_weighted_nobaseline_fit_results')
            if os.path.exists(gaussian_folder):
                gaussian_file = os.path.join(gaussian_folder, 'fit_results.xlsx')
                if os.path.exists(gaussian_file):
                    try:
                        df_chi = pd.read_excel(gaussian_file, sheet_name='Chi_Square_Statistics')
                        global_row = df_chi[df_chi.iloc[:, 0] == 'GLOBAL']
                        if not global_row.empty:
                            sim_data.update({
                                'ratio_calc_gaussian': global_row['ratio'].values[0],
                                'ratio_calc_gaussian_err': global_row['ratio_err'].values[0],
                                'real_calc_diff_over_err_gaussian': abs(global_row['ratio'].values[0] - ratio_real) / global_row['ratio_err'].values[0]
                            })
                    except Exception as e:
                        print(f"Error reading gaussian file {gaussian_file}: {e}")

            # Process ratio-constrained VOIGT fit
            ratio_folder_voigt = os.path.join(simulation_path, 'ratio_sym_voigt_weighted_nobaseline_fit_results')
            if os.path.exists(ratio_folder_voigt):
                ratio_file_voigt = os.path.join(ratio_folder_voigt, 'fit_results.xlsx')
                if os.path.exists(ratio_file_voigt):
                    try:
                        df_chi = pd.read_excel(ratio_file_voigt, sheet_name='Chi_Square_Statistics')
                        global_row = df_chi[df_chi['Line_ID'] == 'GLOBAL']
                        if not global_row.empty:
                            ratio_val = global_row['ratio'].values[0]
                            ratio_err = global_row['ratio_err'].values[0]
                            if pd.notnull(ratio_val) and pd.notnull(ratio_err):
                                sim_data.update({
                                    'ratio_constrained_ratio': ratio_val,
                                    'ratio_constrained_ratio_err': ratio_err,
                                    'real_calc_diff_over_err_ratio_constrained': abs(ratio_val - ratio_real) / ratio_err if ratio_err != 0 else np.nan
                                })
                                print(f"Found ratio constrained VOIGT param for {simulation_folder}: {ratio_val} ± {ratio_err}")
                    except Exception as e:
                        print(f"Error reading ratio file {ratio_file_voigt}: {e}")

            # ----------------------------------------------------------------
            # Process ratio-constrained GAUSSIAN fit # <-- NEW
            # ----------------------------------------------------------------
            ratio_folder_gauss = os.path.join(simulation_path, 'ratio_sym_gaussian_weighted_nobaseline_fit_results')
            if os.path.exists(ratio_folder_gauss):
                ratio_file_gauss = os.path.join(ratio_folder_gauss, 'fit_results.xlsx')
                if os.path.exists(ratio_file_gauss):
                    try:
                        df_chi = pd.read_excel(ratio_file_gauss, sheet_name='Chi_Square_Statistics')
                        global_row = df_chi[df_chi['Line_ID'] == 'GLOBAL']
                        if not global_row.empty:
                            ratio_val = global_row['ratio'].values[0]
                            ratio_err = global_row['ratio_err'].values[0]
                            if pd.notnull(ratio_val) and pd.notnull(ratio_err):
                                sim_data.update({
                                    'ratio_constrained_gaussian_ratio': ratio_val,  # <-- NEW key
                                    'ratio_constrained_gaussian_ratio_err': ratio_err,  # <-- NEW
                                    'real_calc_diff_over_err_ratio_constrained_gaussian': abs(ratio_val - ratio_real) / ratio_err if ratio_err != 0 else np.nan  # <-- NEW
                                })
                                print(f"Found ratio constrained GAUSSIAN param for {simulation_folder}: {ratio_val} ± {ratio_err}")
                    except Exception as e:
                        print(f"Error reading ratio (gaussian) file {ratio_file_gauss}: {e}")

            data.append(sim_data)

        # Once we processed all simulation_folders in the parent, build a DF
        if data:
            df_simulation = pd.DataFrame(data)
            print(f"\nColumns in DataFrame: {df_simulation.columns.tolist()}")

            # Calculate averages for each method
            summary_row = {'simulation': 'AVERAGE', 'ratio_real': ratio_real}

            # 1) VOIGT average
            if 'ratio_calc_voigt' in df_simulation.columns:
                errors = df_simulation['ratio_calc_voigt_err'].dropna()
                ratios = df_simulation['ratio_calc_voigt'].dropna()
                if len(errors) > 0 and len(ratios) > 0:
                    # Suppose 'ratios' is an array of all ratio_i
                    # and 'errors' is an array of all err_i (the 1σ for each ratio_i)
                    inv_var = 1.0 / (errors ** 2)  # 1 / σ_i^2
                    sum_w = inv_var.sum()  # total weight
                    sum_wr = (ratios * inv_var).sum()  # weighted sum
                    avg_ratio_weighted = sum_wr / sum_w
                    avg_ratio_weighted_err = np.sqrt(1.0 / sum_w)

                    avg_ratio = avg_ratio_weighted
                    avg_error = avg_ratio_weighted_err
                    avg_uncertainty = abs(avg_ratio - ratio_real) / avg_error

                    avg_uncertainty = abs(avg_ratio - ratio_real) / avg_error
                    summary_row.update({
                        'ratio_calc_voigt': avg_ratio,
                        'ratio_calc_voigt_err': avg_error,
                        'real_calc_diff_over_err_voigt': avg_uncertainty
                    })

            # 2) GAUSSIAN average
            if 'ratio_calc_gaussian' in df_simulation.columns:
                errors = df_simulation['ratio_calc_gaussian_err'].dropna()
                ratios = df_simulation['ratio_calc_gaussian'].dropna()
                if len(errors) > 0 and len(ratios) > 0:
                    avg_ratio = ratios.mean()
                    avg_error = errors.mean() / np.sqrt(len(errors))
                    avg_uncertainty = abs(avg_ratio - ratio_real) / avg_error
                    summary_row.update({
                        'ratio_calc_gaussian': avg_ratio,
                        'ratio_calc_gaussian_err': avg_error,
                        'real_calc_diff_over_err_gaussian': avg_uncertainty
                    })

            # 3) Ratio-constrained VOIGT average
            if 'ratio_constrained_ratio' in df_simulation.columns:
                errors = df_simulation['ratio_constrained_ratio_err'].dropna()
                ratios = df_simulation['ratio_constrained_ratio'].dropna()
                if len(errors) > 0 and len(ratios) > 0:
                    avg_ratio = ratios.mean()
                    avg_error = errors.mean() / np.sqrt(len(errors))
                    avg_uncertainty = abs(avg_ratio - ratio_real) / avg_error
                    summary_row.update({
                        'ratio_constrained_ratio': avg_ratio,
                        'ratio_constrained_ratio_err': avg_error,
                        'real_calc_diff_over_err_ratio_constrained': avg_uncertainty
                    })

            # ----------------------------------------------------------------
            # 4) Ratio-constrained GAUSSIAN average # <-- NEW
            # ----------------------------------------------------------------
            if 'ratio_constrained_gaussian_ratio' in df_simulation.columns:
                errors = df_simulation['ratio_constrained_gaussian_ratio_err'].dropna()
                ratios = df_simulation['ratio_constrained_gaussian_ratio'].dropna()
                if len(errors) > 0 and len(ratios) > 0:
                    avg_ratio = ratios.mean()
                    avg_error = errors.mean() / np.sqrt(len(errors))
                    avg_uncertainty = abs(avg_ratio - ratio_real) / avg_error
                    summary_row.update({
                        'ratio_constrained_gaussian_ratio': avg_ratio,
                        'ratio_constrained_gaussian_ratio_err': avg_error,
                        'real_calc_diff_over_err_ratio_constrained_gaussian': avg_uncertainty
                    })

            # Add summary row
            df_simulation = pd.concat([df_simulation, pd.DataFrame([summary_row])], ignore_index=True)

            # Save to Excel with highlighting
            output_file = os.path.join(parent_root, f'{parent_folder}_summary.xlsx')
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                df_simulation.to_excel(writer, index=False)

                # Get the worksheet
                worksheet = writer.sheets['Sheet1']

                # Find the columns to highlight
                cols_to_highlight = {}
                for idx, col in enumerate(df_simulation.columns, 1):
                    if col.startswith('real_calc_diff_over_err_'):
                        cols_to_highlight[col] = get_column_letter(idx)

                # Find the row with 'AVERAGE'
                for row_idx, cell in enumerate(worksheet['A'], 1):
                    if cell.value == 'AVERAGE':
                        # Highlight the cells
                        yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
                        for col_letter in cols_to_highlight.values():
                            worksheet[f'{col_letter}{row_idx}'].fill = yellow_fill
                        break

            print(f"Saved parent folder summary to {output_file}")


if __name__ == '__main__':
    base_path = "/Users/tomsayada/spectral_analysis_project/data/"
    process_simulation_folder(base_path)
